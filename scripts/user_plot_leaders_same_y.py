import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import openpyxl as xl

from openpyxl import load_workbook

workbook = load_workbook(filename="/Users/gbtmdr/Desktop/CNEO_CPT_ESI_data_with_dE_NH_heatmap_desktop.xlsx")

#------------------------
# Variables and Functions
#------------------------
sheet_hf=workbook["Table S2 (HF)n barriers"]
sheet_hcl=workbook["Table S1 (HCl)n barriers"]
sheet_water=workbook["Table S3 (H2O)n barriers"]

mon_units = [3, 4, 5]
columns_dft = [3, 10, 17]
columns_cneo = [6, 13, 20]
functionals_qz = [cell[0].value for cell in sheet_hf['A7:A25']]

hf_refs = [20.68, 14.80, 14.84]
hcl_refs = [27.35, 30.88, 37.90]
water_refs = [29.99, 26.88, 30.41]

# X-axis distinct categories
monomers = np.array([3, 4, 5])
num_functionals = len(functionals_qz)

# Visual assets for line distinctions
colors = plt.cm.tab20(np.linspace(0, 1, num_functionals))
markers = ['o', 's', '^', 'D', 'v', '<', '>', 'p', '*', 'h']

# Leader-line labels: (y at the n=5 end, functional name, line colour, marker) collected per panel
label_items = {}

def spread(ys, lo, hi, gap):
    # push sorted label positions apart until neighbours are at least `gap` apart, staying inside [lo, hi]
    pos = list(ys)
    for _ in range(200):
        moved = False
        for i in range(1, len(pos)):
            if pos[i] - pos[i - 1] < gap:
                shift = (gap - (pos[i] - pos[i - 1])) / 2
                pos[i - 1] -= shift
                pos[i] += shift
                moved = True
        # labels pushed past the bottom (top) edge are moved back inside, carrying their crowded neighbours with them
        pos[0] = max(pos[0], lo)
        for i in range(1, len(pos)):
            pos[i] = max(pos[i], pos[i - 1] + gap)
        pos[-1] = min(pos[-1], hi)
        for i in range(len(pos) - 2, -1, -1):
            pos[i] = min(pos[i], pos[i + 1] - gap)
        if not moved:
            break
    return pos

def label_lines(fig, ax, items, x_line=5.02, x_text=5.32, fontsize=7, gap_pts=8.5):
    # write each functional name just right of the panel, with its marker in front and a leader line back to its line end
    fig.canvas.draw()
    height_pts = ax.get_window_extent().height * 72 / fig.dpi
    lo, hi = ax.get_ylim()
    ax.set_ylim(lo, hi)   # freeze the range so the label markers below do not rescale the axis
    gap = (hi - lo) * gap_pts / height_pts
    items = sorted(items, key=lambda t: t[0])
    ys = spread([t[0] for t in items], lo + gap / 2, hi - gap / 2, gap)
    for (y_line, text, colour, marker), y_text in zip(items, ys):
        ax.plot(x_text, y_text, marker=marker, color=colour, linestyle="", markersize=5, clip_on=False)
        ax.annotate(text, xy=(x_text, y_text), xytext=(7, 0), textcoords="offset points",
                    fontsize=fontsize, ha="left", va="center", annotation_clip=False)
        ax.annotate("", xy=(x_line, y_line), xytext=(x_text, y_text), textcoords="data", annotation_clip=False,
                    arrowprops=dict(arrowstyle="-", color=colour, linewidth=0.7, shrinkA=4, shrinkB=2))


#------------------------
# Trimer Data
#------------------------

dft_barriers_hf_trimer_qz = sheet_hf['C7:C25']
#dft_barriers_hf_trimer_tz
#dft_barriers_hf_trimer_sv

dft_barriers_hcl_trimer_qz = sheet_hcl['C7:C25']
#dft_barriers_hcl_trimer_tz
#dft_barriers_hcl_trimer_sv

dft_barriers_water_trimer_qz = sheet_water['C7:C25']
#dft_barriers_water_trimer_tz
#dft_barriers_water_trimer_sv

cneo_barriers_hf_trimer_qz = sheet_hf['F7:F25']
#cneo_barriers_hf_trimer_tz
#cneo_barriers_hf_trimer_sv

cneo_barriers_hcl_trimer_qz = sheet_hcl['F7:F25']
#cneo_barriers_hcl_trimer_tz
#cneo_barriers_hcl_trimer_sv

cneo_barriers_water_trimer_qz = sheet_water['F7:F25']
#cneo_barriers_water_trimer_tz
#cneo_barriers_water_trimer_sv

#------------------------
# Tetramer Data
#------------------------

dft_barriers_hf_tetramer_qz = sheet_hf['J7:J25']
#dft_barriers_hf_tetramer_tz
#dft_barriers_hf_tetramer_sv

dft_barriers_hcl_tetramer_qz = sheet_hcl['J7:J25']
#dft_barriers_hcl_tetramer_tz
#dft_barriers_hcl_tetramer_sv

dft_barriers_water_tetramer_qz = sheet_water['J7:J25']
#dft_barriers_water_tetramer_tz
#dft_barriers_water_tetramer_sv

cneo_barriers_hf_tetramer_qz = sheet_hf['M7:M25']
#cneo_barriers_hf_tetramer_tz
#cneo_barriers_hf_tetramer_sv

cneo_barriers_hcl_tetramer_qz = sheet_hcl['M7:M25']
#cneo_barriers_hcl_tetramer_tz
#cneo_barriers_hcl_tetramer_sv

cneo_barriers_water_tetramer_qz = sheet_water['M7:M25']
#cneo_barriers_water_tetramer_tz
#cneo_barriers_water_tetramer_sv

#------------------------
# Pentamer Data
#------------------------

dft_barriers_hf_pentamer_qz = sheet_hf['Q7:Q25']
#dft_barriers_hf_pentamer_tz
#dft_barriers_hf_pentamer_sv

dft_barriers_hcl_pentamer_qz = sheet_hcl['Q7:Q25']
#dft_barriers_hcl_pentamer_tz
#dft_barriers_hcl_pentamer_sv

dft_barriers_water_pentamer_qz = sheet_water['Q7:Q25']
#dft_barriers_water_pentamer_tz
#dft_barriers_water_pentamer_sv

cneo_barriers_hf_pentamer_qz = sheet_hf['T7:T25']
#cneo_barriers_hf_pentamer_tz
#cneo_barriers_hf_pentamer_sv

cneo_barriers_hcl_pentamer_qz = sheet_hcl['T7:T25']
#cneo_barriers_hcl_pentamer_tz
#cneo_barriers_hcl_pentamer_sv

cneo_barriers_water_pentamer_qz = sheet_water['T7:T25']
#cneo_barriers_water_pentamer_tz
#cneo_barriers_water_pentamer_sv



fig, axes = plt.subplots(nrows=3, ncols=2, figsize=(10,12))

grid_map = {

    (0,0): {"title": "(HF)$_{n}$ DFT Barriers", "sheet_name": "Table S2 (HF)n barriers", "col3_idx":3, "col4_idx":10, "col5_idx":17, "refs": hf_refs, "method": "DFT"},
    (0,1): {"title": "(HF)$_{n}$ CNEO-DFT Barriers", "sheet_name": "Table S2 (HF)n barriers", "col3_idx":6, "col4_idx":13, "col5_idx":20, "refs": hf_refs, "method": "CNEO-DFT"},
    (1,0): {"title": "(HCl)$_{n}$ DFT Barriers", "sheet_name": "Table S1 (HCl)n barriers", "col3_idx":3, "col4_idx":10, "col5_idx":17, "refs": hcl_refs, "method": "DFT"},
    (1,1): {"title": "(HCl)$_{n}$ CNEO-DFT Barriers", "sheet_name": "Table S1 (HCl)n barriers", "col3_idx":6, "col4_idx":13, "col5_idx":20, "refs": hcl_refs, "method": "CNEO-DFT"},
    (2,0): {"title": "(H$_{2}$O)$_{n}$ DFT Barriers", "sheet_name": "Table S3 (H2O)n barriers", "col3_idx":3, "col4_idx":10, "col5_idx":17, "refs": water_refs, "method": "DFT"},
    (2,1): {"title": "(H$_{2}$O)$_{n}$ CNEO-DFT Barriers", "sheet_name": "Table S3 (H2O)n barriers", "col3_idx":6, "col4_idx":13, "col5_idx":20, "refs": water_refs, "method": "CNEO-DFT"}
}


for row in range(3):
    for col in range(2):

        ax = axes[row, col]

        config = grid_map [(row, col)]
        ax.set_title(config["title"], fontsize=12, fontweight='bold')

        worksheet = workbook[config["sheet_name"]]

        current_row = 7
        label_items[(row, col)] = []
        panel_energies = []

        for func_idx, func_name in enumerate(functionals_qz):

            val_3 = worksheet.cell(row=current_row, column=config["col3_idx"]).value
            val_4 = worksheet.cell(row=current_row, column=config["col4_idx"]).value
            val_5 = worksheet.cell(row=current_row, column=config["col5_idx"]).value
            y_energies = [val_3, val_4, val_5]
            panel_energies.append(y_energies)

            current_row += 1

            ax.plot(
                monomers,
                y_energies,
                label=func_name,
                color=colors[func_idx],
                marker=markers[func_idx % len(markers)],
                linewidth=1.5,
                markersize=6
            )

            label_items[(row, col)].append((val_5, func_name, colors[func_idx], markers[func_idx % len(markers)]))

        # Mean over the DFT functionals (absent barriers, entered as 0.00 in the tables, are left out)
        mean_energies = [np.mean([y[i] for y in panel_energies if y[i] > 0]) for i in range(3)]

        ax.plot(
            monomers,
            mean_energies,
            label="Mean " + config["method"],
            color="black",
            linestyle="--",
            marker="D",
            linewidth=3,
            markersize=7
        )

        label_items[(row, col)].append((mean_energies[2], "Mean " + config["method"], "black", "D"))

        # CCSD(T)-F12 reference barriers in black
        ax.plot(
            monomers,
            config["refs"],
            label="CCSD(T)-F12",
            color="black",
            marker="s",
            linewidth=3,
            markersize=7
        )

        label_items[(row, col)].append((config["refs"][2], "CCSD(T)-F12", "black", "s"))



        # Standard individual subplot visual formatting configurations
        ax.set_xlabel("Monomer Units")
        ax.set_ylabel("Barrier Energy (kcal/mol)")
        ax.set_xticks(monomers)
        ax.set_xlim(2.8, 5.2)
        ax.grid(True, linestyle="--", alpha=0.4)

fig.suptitle("Barriers for Concerted Proton Transfer with respect to Ring Size", fontsize=16, fontweight="bold")

# Fixed layout with a gutter to the right of every panel for the functional names
fig.subplots_adjust(left=0.06, right=0.86, top=0.94, bottom=0.05, hspace=0.35, wspace=0.55)

# Same y-axis range on every panel
y_lo = min(ax.get_ylim()[0] for ax in axes.flat)
y_hi = max(ax.get_ylim()[1] for ax in axes.flat)
for ax in axes.flat:
    ax.set_ylim(y_lo, y_hi)

# Functional names with leader lines, placed once the layout is known
for (row, col), items in label_items.items():
    label_lines(fig, axes[row, col], items)

plt.show()
